"""
Step 4: Build the Excel workbook with live formulas (SUMIFS, INDEX/MATCH, VLOOKUP),
KPI cards, summary tables and native charts.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

df = pd.read_csv("data/clean/sales_clean.csv", parse_dates=["order_date"])
products = pd.read_csv("data/raw/products.csv")
regions = pd.read_csv("data/raw/regions.csv")
N = len(df)
LAST = N + 1

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
KPI_FILL = PatternFill("solid", fgColor="EAF2FB")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------- Lookups ----------------
wp = wb.active
wp.title = "Products"
wp.append(list(products.columns))
for r in products.itertuples(index=False):
    wp.append(list(r))
style_header(wp, 1, products.shape[1])
autosize(wp, [12, 30, 18, 12, 12])

wr = wb.create_sheet("Regions")
wr.append(list(regions.columns))
for r in regions.itertuples(index=False):
    wr.append(list(r))
style_header(wr, 1, regions.shape[1])
autosize(wr, [10, 12, 20])

# ---------------- Data (with lookup formula columns) ----------------
wd = wb.create_sheet("Data")
cols = ["order_id", "order_date", "customer_id", "segment", "region", "state", "product_id", "quantity",
        "discount", "sales", "profit", "ship_mode", "product_name", "category", "year", "month",
        "month_name", "year_month", "quarter", "profit_margin_pct"]
extra = ["regional_manager", "unit_price", "list_revenue", "discount_given"]
wd.append(cols + extra)
for i, row in enumerate(df[cols].itertuples(index=False), start=2):
    vals = list(row)
    vals[1] = vals[1].to_pydatetime()
    wd.append(vals + [
        f"=INDEX(Regions!$C$2:$C$5,MATCH(E{i},Regions!$B$2:$B$5,0))",
        f"=VLOOKUP(G{i},Products!$A$2:$E$25,4,FALSE)",
        f"=V{i}*H{i}",
        f"=W{i}-J{i}",
    ])
    wd.cell(row=i, column=2).number_format = "yyyy-mm-dd"
style_header(wd, 1, len(cols) + len(extra))
wd.freeze_panes = "A2"
wd.auto_filter.ref = f"A1:{get_column_letter(len(cols) + len(extra))}{LAST}"
autosize(wd, [13, 12, 12, 13, 9, 15, 11, 9, 9, 11, 11, 15, 26, 16, 7, 7, 10, 11, 9, 12, 16, 11, 13, 13])

R = lambda col: f"Data!${col}$2:${col}${LAST}"   # absolute range helper
SALES, PROFIT, REGION, CAT, PROD, YM, QTY, CUST, ORD, SEG, DISC = (
    R("J"), R("K"), R("E"), R("N"), R("M"), R("R"), R("H"), R("C"), R("A"), R("D"), R("I"))

# ---------------- KPI sheet ----------------
wk = wb.create_sheet("KPI Summary", 0)
wk["A1"] = "Sales Performance - KPI Summary (Jan 2024 - Dec 2025)"
wk["A1"].font = Font(size=16, bold=True, color="1F3A5F")
kpis = [
    ("Total Sales", f"=SUM({SALES})", '#,##0'),
    ("Total Profit", f"=SUM({PROFIT})", '#,##0'),
    ("Profit Margin %", f"=B4/B3", '0.0%'),
    ("Total Orders", f"=COUNTA({ORD})", '#,##0'),
    ("Avg Order Value", f"=B3/B6", '#,##0'),
    ("Units Sold", f"=SUM({QTY})", '#,##0'),
    ("Total Discount Given", f"=SUM(Data!$X$2:$X${LAST})", '#,##0'),
    ("Avg Discount Rate", f"=AVERAGE({DISC})", '0.0%'),
]
wk.append([])
wk["A2"], wk["B2"] = "KPI", "Value"
style_header(wk, 2, 2)
for i, (name, f, fmt) in enumerate(kpis, start=3):
    wk.cell(row=i, column=1, value=name).font = Font(bold=True)
    c = wk.cell(row=i, column=2, value=f)
    c.number_format = fmt
    for col in (1, 2):
        wk.cell(row=i, column=col).fill = KPI_FILL
        wk.cell(row=i, column=col).border = BORDER
autosize(wk, [26, 18, 4, 18, 14, 14, 14, 14])

# ---------------- Region summary (SUMIFS + manager via INDEX/MATCH) ----------------
ws = wb.create_sheet("Region Summary")
ws.append(["Region", "Manager", "Sales", "Profit", "Margin %", "% of Total Sales", "Orders"])
style_header(ws, 1, 7)
for i, reg in enumerate(["West", "South", "North", "East"], start=2):
    ws.append([reg,
               f"=INDEX(Regions!$C$2:$C$5,MATCH(A{i},Regions!$B$2:$B$5,0))",
               f"=SUMIFS({SALES},{REGION},A{i})",
               f"=SUMIFS({PROFIT},{REGION},A{i})",
               f"=D{i}/C{i}",
               f"=C{i}/SUM($C$2:$C$5)",
               f"=COUNTIFS({REGION},A{i})"])
    for col, fmt in ((3, "#,##0"), (4, "#,##0"), (5, "0.0%"), (6, "0.0%"), (7, "#,##0")):
        ws.cell(row=i, column=col).number_format = fmt
autosize(ws, [12, 18, 14, 14, 11, 16, 10])
ch = BarChart(); ch.type = "col"; ch.title = "Sales & Profit by Region"
ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=1, max_row=5), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
ch.height, ch.width = 8, 14
ws.add_chart(ch, "I2")

# ---------------- Category summary ----------------
wc = wb.create_sheet("Category Summary")
wc.append(["Category", "Sales", "Profit", "Margin %"])
style_header(wc, 1, 4)
for i, cat in enumerate(["Technology", "Furniture", "Office Supplies"], start=2):
    wc.append([cat, f"=SUMIFS({SALES},{CAT},A{i})", f"=SUMIFS({PROFIT},{CAT},A{i})", f"=C{i}/B{i}"])
    for col, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "0.0%")):
        wc.cell(row=i, column=col).number_format = fmt
autosize(wc, [18, 14, 14, 11])
ch = BarChart(); ch.type = "bar"; ch.title = "Sales by Category"
ch.add_data(Reference(wc, min_col=2, min_row=1, max_row=4), titles_from_data=True)
ch.set_categories(Reference(wc, min_col=1, min_row=2, max_row=4))
ch.height, ch.width = 7, 13
wc.add_chart(ch, "G2")

# ---------------- Product summary ----------------
wpd = wb.create_sheet("Product Summary")
wpd.append(["Product", "Category (lookup)", "Units", "Sales", "Profit", "Margin %", "Sales Rank"])
style_header(wpd, 1, 7)
for i, name in enumerate(products["product_name"], start=2):
    n_last = len(products) + 1
    wpd.append([name,
                f"=INDEX(Products!$C$2:$C$25,MATCH(A{i},Products!$B$2:$B$25,0))",
                f"=SUMIFS({QTY},{PROD},A{i})",
                f"=SUMIFS({SALES},{PROD},A{i})",
                f"=SUMIFS({PROFIT},{PROD},A{i})",
                f"=E{i}/D{i}",
                f"=RANK(D{i},$D$2:$D${n_last})"])
    for col, fmt in ((3, "#,##0"), (4, "#,##0"), (5, "#,##0"), (6, "0.0%")):
        wpd.cell(row=i, column=col).number_format = fmt
autosize(wpd, [30, 18, 9, 14, 14, 11, 11])
wpd.auto_filter.ref = f"A1:G{len(products) + 1}"
ch = BarChart(); ch.type = "bar"; ch.title = "Sales by Product"
ch.add_data(Reference(wpd, min_col=4, min_row=1, max_row=len(products) + 1), titles_from_data=True)
ch.set_categories(Reference(wpd, min_col=1, min_row=2, max_row=len(products) + 1))
ch.height, ch.width = 14, 16
wpd.add_chart(ch, "I2")

# ---------------- Monthly trend ----------------
wm = wb.create_sheet("Monthly Trend")
wm.append(["Month", "Sales", "Profit", "MoM Growth %"])
style_header(wm, 1, 4)
months = sorted(df["year_month"].unique())
for i, m in enumerate(months, start=2):
    wm.append([m, f"=SUMIFS({SALES},{YM},A{i})", f"=SUMIFS({PROFIT},{YM},A{i})",
               "" if i == 2 else f"=B{i}/B{i - 1}-1"])
    for col, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "0.0%")):
        wm.cell(row=i, column=col).number_format = fmt
autosize(wm, [11, 14, 14, 14])
ch = LineChart(); ch.title = "Monthly Sales & Profit"
ch.add_data(Reference(wm, min_col=2, max_col=3, min_row=1, max_row=len(months) + 1), titles_from_data=True)
ch.set_categories(Reference(wm, min_col=1, min_row=2, max_row=len(months) + 1))
ch.height, ch.width = 8, 18
wm.add_chart(ch, "F2")

# ---------------- Discount impact ----------------
wdsc = wb.create_sheet("Discount Impact")
wdsc.append(["Discount Level", "Sales", "Profit", "Margin %"])
style_header(wdsc, 1, 4)
for i, d in enumerate([0, 0.05, 0.10, 0.15, 0.20, 0.30], start=2):
    wdsc.append([d, f"=SUMIFS({SALES},{DISC},A{i})", f"=SUMIFS({PROFIT},{DISC},A{i})", f"=C{i}/B{i}"])
    for col, fmt in ((1, "0%"), (2, "#,##0"), (3, "#,##0"), (4, "0.0%")):
        wdsc.cell(row=i, column=col).number_format = fmt
autosize(wdsc, [14, 14, 14, 11])
ch = BarChart(); ch.type = "col"; ch.title = "Profit Margin % by Discount Level"
ch.add_data(Reference(wdsc, min_col=4, min_row=1, max_row=7), titles_from_data=True)
ch.set_categories(Reference(wdsc, min_col=1, min_row=2, max_row=7))
ch.height, ch.width = 7, 13
wdsc.add_chart(ch, "G2")

wb.move_sheet("Data", offset=len(wb.sheetnames))  # Data last
wb.save("outputs/Sales_Analysis.xlsx")
print("Saved outputs/Sales_Analysis.xlsx")
